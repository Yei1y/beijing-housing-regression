"""
脚本：重构报告结构——将模型设定移到 EDA 之后
==============================================
操作：
1. 从第二节移除模型设定（原 2.3）
2. 重命名第二节子标题（2.4→2.3, 2.5→2.4）
3. 在 EDA 后插入新的"四、模型设定与变量选择"
4. 原"四、变量选择"降级为其中一节
"""

import sys
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt


def insert_paragraph_after(doc, ref_para, text, style_name='Normal'):
    """在 ref_para 后插入一个新段落。"""
    new_p = OxmlElement('w:p')
    ref_para._element.addnext(new_p)
    new_para = Paragraph(new_p, ref_para._parent)
    run = new_para.add_run(text)
    try:
        new_para.style = doc.styles[style_name]
    except Exception:
        pass
    return new_para


def set_para_text(para, text):
    """清空段落文本并写入新文本。"""
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)


def clear_para(para):
    """完全清空段落。"""
    for run in para.runs:
        run.text = ""


def set_para_style(para, style_name):
    """设置段落样式。"""
    try:
        para.style = para._parent.styles[style_name]
    except Exception:
        pass


def main():
    doc = Document("report/分析报告.docx")
    paras = doc.paragraphs

    # ===== 1. 改造第二节：移除 2.3 模型设定 =====
    print("1. 改造第二节...")

    # 改节标题: "二、数据与方法论" → "二、数据来源与预处理"
    for i, p in enumerate(paras):
        if p.text.strip().startswith("二、数据"):
            set_para_text(p, "二、数据来源与预处理")
            print(f"   节标题已改 (para {i})")
            break

    # 清空 2.3 模型设定的所有段落 (para 32-41)
    # 2.3 标题在 para 32
    clear_para(paras[32])  # "2.3 模型设定"
    for i in range(33, 42):
        clear_para(paras[i])
    print("   2.3 模型设定已清空")

    # 重命名子标题
    # "2.4 数据预处理" → "2.3 数据预处理"
    for i, p in enumerate(paras):
        if p.text.strip().startswith("2.4 数据预处理"):
            set_para_text(p, "2.3 数据预处理")
            print(f"   重命名 para {i}: 2.4→2.3")
            break

    # "2.4.1 关于DOM变量缺失值的讨论" → "2.3.1"
    for i, p in enumerate(paras):
        if p.text.strip().startswith("2.4.1"):
            # Keep the title but change numbering
            text = p.text.strip()
            new_text = text.replace("2.4.1", "2.3.1", 1)
            set_para_text(p, new_text)
            print(f"   重命名 para {i}: 2.4.1→2.3.1")
            break

    # "2.4.2 共线性问题的预设考虑" → "2.3.2"
    for i, p in enumerate(paras):
        if p.text.strip().startswith("2.4.2"):
            text = p.text.strip()
            new_text = text.replace("2.4.2", "2.3.2", 1)
            set_para_text(p, new_text)
            print(f"   重命名 para {i}: 2.4.2→2.3.2")
            break

    # "2.5 训练-测试集划分与标准化" → "2.4 训练-测试集划分与标准化"
    for i, p in enumerate(paras):
        if p.text.strip().startswith("2.5"):
            text = p.text.strip()
            new_text = text.replace("2.5", "2.4", 1)
            set_para_text(p, new_text)
            print(f"   重命名 para {i}: 2.5→2.4")
            break

    # ===== 2. 找到 EDA 结束位置 =====
    print("2. 定位 EDA 结束位置...")
    eda_end_idx = None
    for i, p in enumerate(paras):
        if p.text.strip().startswith("四、变量选择"):
            eda_end_idx = i - 1  # 前一段是 EDA 的最后一段
            print(f"   EDA 结束于 para {eda_end_idx}")
            break

    if eda_end_idx is None:
        print("错误：找不到 EDA 结束位置")
        return

    ref_para = paras[eda_end_idx]

    # ===== 3. 在 EDA 后插入新节"四、模型设定与变量选择" =====
    print("3. 插入新节 四、模型设定与变量选择...")

    # 获取旧 2.3 的内容（从已修改的版本中提取模型设定的内容）
    # 实际上旧 2.3 已经被清空了，我需要重新写模型设定的内容
    # 这些内容已经在前一轮修改中被写入清空的段落中
    # 我需要在插入的地方重建模型设定内容

    # 插入新节标题
    h1 = insert_paragraph_after(doc, ref_para, "四、模型设定与变量选择", 'Heading 1')
    print("   插入节标题: 四、模型设定与变量选择")

    # 插入 4.1 模型设定
    h2_1 = insert_paragraph_after(doc, h1, "4.1 基准回归模型", 'Heading 2')
    insert_paragraph_after(doc, h2_1,
        "本研究在特征价格模型框架下，设定包含时间和区域双向固定效应的半对数回归方程作为基准模型：")
    insert_paragraph_after(doc, h2_1,
        "log(Price_{it}) = β₀ + X_{it}′β + γ_t + δ_k + ε_{it}")
    insert_paragraph_after(doc, h2_1,
        "其中 log(Price_{it}) 为第 i 条交易在第 t 期的对数价格；"
        "X_{it} 包含连续型房屋和区位特征（dist_center、Lat、Lng、square 等）；"
        "γ_t 为时间固定效应，由年度虚拟变量（trade_year_2012 至 trade_year_2017，以 2011 年为基期）"
        "和月度虚拟变量（trade_month_2 至 trade_month_12，以 1 月为基期）共同实现；"
        "δ_k 为区域固定效应。")
    insert_paragraph_after(doc, h2_1,
        "该设定的关键优势在于不施加「年度价格增幅恒定」的线性约束，"
        "允许数据自行估计每年的基准价格水平。"
        "事实上，EDA 部分（第三节）已表明 2011-2017 年房价趋势呈明显非线性特征——"
        "2015 年前平缓上涨、2015 年后加速攀升，因此 γ_t 的灵活设定远优于线性趋势变量。"
        "月度虚拟变量则进一步控制季节性波动（如年初和年末的交易量差异）。")

    print("   4.1 模型设定已插入")

    # ===== 4. 改造原"四、变量选择" =====
    print("4. 改造原四、变量选择...")

    # 从 Heading1 改为 Heading2，文本改为"4.2 变量选择的共识策略"
    for i, p in enumerate(paras):
        if p.text.strip().startswith("四、变量选择"):
            set_para_text(p, "4.2 变量选择的共识策略")
            set_para_style(p, 'Heading 2')
            print(f"   改造 para {i}: 四、变量选择 → 4.2 变量选择的共识策略")

        # 原 4.1 LASSO正则化 → 4.2.1
        if p.text.strip().startswith("4.1 LASSO"):
            set_para_text(p, "4.2.1 LASSO 正则化与惩罚力度问题")
            print(f"   改造 para {i}: 4.1→4.2.1")

    # ===== 5. 更新后续章节标题（在原结构基础上） =====
    # 原五→五、原六→六、原七→七 无需改号，因为新四插入后四变为了模型设定
    # 但原四(变量选择)被吸收进了新四，所以原五→还是五
    # 但需要检查"五、模型诊断"是否还是正确的
    print("5. 验证后续章节标题...")
    for i, p in enumerate(paras):
        t = p.text.strip()
        if t.startswith("五、") and '诊断' in t:
            print(f"   保留: para {i} = {t[:30]}")
        elif t.startswith("六、") and '比较' in t:
            print(f"   保留: para {i} = {t[:30]}")
        elif t.startswith("七、") and '结论' in t:
            print(f"   保留: para {i} = {t[:30]}")

    # ===== 6. 保存 =====
    out_path = "report/分析报告.docx"
    doc.save(out_path)
    print(f"\n✅ 已保存: {out_path}")


if __name__ == "__main__":
    main()
